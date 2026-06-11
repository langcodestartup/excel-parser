"""Multi-table extraction goldens — Phase 10b (plan v2 §4 Task 10.2).

Layers covered:

1.  **Step 5 goldens** — both stacked tables of ``multi_table_stacked`` are
    extracted as independent :class:`TableResult` items (T1/T2, top-down).
2.  **Guard 1** — ``stacked_uneven_width``: the narrow 3-column table is not
    diluted by the sheet-global ``max_col`` (8) denominator.
3.  **Guard 5** — the *second* (lower) block's ReadPlan is pinned by an actual
    ``pandas.read_excel`` round-trip via ``adapters.pandas_loader``, so the
    "rows above the header are absorbed wholesale" rule is measured, not
    assumed.
4.  **Guard 4** — ``SheetOverride.header_row`` / ``skip_rows_add`` are
    absolute 1-based sheet coordinates applied only to the block containing
    the row (block-2-targeted override tests).
5.  **Guard 6/7 + mirror** — deterministic warning order, blocks hold only
    table-judged bands, the flat fields mirror ``blocks[0]``, and the mirror
    block's plan equals ``sheet.read_plan`` across the whole corpus.

Coordinates in assertions are openpyxl 1-based for ``SheetProfile`` /
``TableBlock`` and pandas 0-based for ``ReadPlan`` fields [D1]; the fixture
layouts are documented in ``tests/fixtures/generate.py`` (single source).
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from excel_inspector import (
    InspectionOptions,
    SheetOverride,
    extract,
    inspect,
)
from excel_inspector.adapters.pandas_loader import load_dataframe


# ---------------------------------------------------------------------------
# Step 5 goldens — both stacked tables extracted (plan v2 §4 Task 10.2)
# ---------------------------------------------------------------------------


def test_two_stacked_tables_both_extracted(fixture_path) -> None:
    wr = extract(fixture_path("multi_table_stacked"))
    ids = [t.table_id for t in wr.tables]
    assert len(ids) == 2 and ids[0].endswith("!T1") and ids[1].endswith("!T2")
    t1, t2 = wr.tables
    assert list(t1.dataframe.columns) == ["부서", "인원", "예산"]
    assert len(t1.dataframe) == 3            # 조용한 누락 제거 — 표1 복원!
    assert list(t2.dataframe.columns) == ["제품명", "단가", "재고", "비고"]
    assert len(t2.dataframe) == 3


def test_stacked_tables_ids_are_top_down(fixture_path) -> None:
    """T1 is the top-most table (header row 1), T2 the lower one (row 7)."""

    wr = extract(fixture_path("multi_table_stacked"))
    t1, t2 = wr.tables
    assert t1.table_id == "Sheet1!T1" and t2.table_id == "Sheet1!T2"
    assert t1.header_row == 1   # 1-based [D1]
    assert t2.header_row == 7
    assert t1.header_provenance == t2.header_provenance == "heuristic"
    assert 0.0 < t1.header_confidence <= 1.0
    assert 0.0 < t2.header_confidence <= 1.0


def test_stacked_tables_values_do_not_leak_across_blocks(fixture_path) -> None:
    """Partial value golden: each table carries only its own band's rows."""

    wr = extract(fixture_path("multi_table_stacked"))
    t1, t2 = wr.tables
    assert list(t1.dataframe["부서"]) == ["영업", "개발", "관리"]
    assert int(t1.dataframe["인원"].sum()) == 37
    assert list(t2.dataframe["제품명"]) == ["키보드", "마우스", "모니터"]
    assert int(t2.dataframe["단가"].sum()) == 255000
    # No cross-band leakage in either direction.
    flat1 = t1.dataframe.astype(str).to_numpy().ravel().tolist()
    assert not any("키보드" in cell for cell in flat1)
    flat2 = t2.dataframe.astype(str).to_numpy().ravel().tolist()
    assert not any("영업" in cell for cell in flat2)


def test_stacked_tables_json_shape(fixture_path) -> None:
    """The JSON 'tables' array carries both blocks (schema v1.0 unchanged)."""

    wr = extract(fixture_path("multi_table_stacked"))
    parsed = json.loads(wr.to_json())
    (sheet,) = parsed["sheets"]
    assert [t["table_id"] for t in sheet["tables"]] == [
        "Sheet1!T1", "Sheet1!T2",
    ]
    t1, t2 = sheet["tables"]
    assert [c["name"] for c in t1["columns"]] == ["부서", "인원", "예산"]
    assert [c["name"] for c in t2["columns"]] == ["제품명", "단가", "재고", "비고"]
    assert t1["row_count"] == 3 and t2["row_count"] == 3


# ---------------------------------------------------------------------------
# Guard 1 — uneven-width stacking: band-local denominator, no dilution
# ---------------------------------------------------------------------------


def test_stacked_uneven_width_narrow_table_recognized(fixture_path) -> None:
    """The narrow 3-col table survives next to the 8-col one (guard 1).

    With a sheet-global ``max_col`` (8) score denominator the narrow band's
    header score would be diluted to ~3/8 of its real value and the band
    misjudged "not a table" — silently re-losing table 1.
    """

    wr = extract(fixture_path("stacked_uneven_width"))
    ids = [t.table_id for t in wr.tables]
    assert ids == ["Sheet1!T1", "Sheet1!T2"]
    t1, t2 = wr.tables
    assert list(t1.dataframe.columns) == ["코드", "명칭", "수량"]
    assert len(t1.dataframe) == 3
    assert list(t2.dataframe.columns) == [
        "일자", "지점", "담당", "품목", "단가", "수량", "금액", "비고",
    ]
    assert len(t2.dataframe) == 3
    assert int(t1.dataframe["수량"].sum()) == 15
    assert int(t2.dataframe["금액"].sum()) == 25000


def test_stacked_uneven_width_narrow_block_boundaries(fixture_path) -> None:
    """Block 1 spans columns A-C only (1-based [D1]); block 2 is full-width."""

    sheet = inspect(fixture_path("stacked_uneven_width")).sheets[0]
    b1, b2 = sheet.blocks
    assert (b1.band_start_row, b1.band_end_row) == (1, 4)
    assert (b1.data_start_row, b1.data_end_row) == (2, 4)
    assert (b1.data_left_col, b1.data_right_col) == (1, 3)
    assert b1.read_plan is not None and b1.read_plan.usecols == "A:C"
    assert (b2.band_start_row, b2.band_end_row) == (7, 10)
    assert (b2.data_start_row, b2.data_end_row) == (8, 10)
    # Full-width table -> no usecols restriction (spec §4.5 convention).
    assert (b2.data_left_col, b2.data_right_col) == (None, None)
    assert b2.read_plan is not None and b2.read_plan.usecols is None


# ---------------------------------------------------------------------------
# Guard 5 — the lower block's ReadPlan pinned by a real pandas round-trip
# ---------------------------------------------------------------------------


def test_block2_read_plan_pandas_round_trip(fixture_path) -> None:
    """블록 2의 ReadPlan을 실제 read_excel 왕복으로 골든 고정 (guard 5).

    "헤더 위 전부 흡수 규칙이 블록 위치와 무관하게 성립"은 가정이 아니라
    실측이어야 한다: 블록 2(헤더 7행)의 plan은 1~6행(표1 + 빈 분리줄)을
    통째로 skiprows에 흡수하고 header를 post-skip 0으로 정규화해야 하며,
    실제 적재 결과가 정확히 3행/4컬럼이어야 한다.
    """

    path = fixture_path("multi_table_stacked")
    sheet = inspect(path).sheets[0]
    assert len(sheet.blocks) == 2
    b2 = sheet.blocks[1]
    plan = b2.read_plan
    assert plan is not None
    # 0-based loading domain [D1]: rows 1-6 (1-based) -> skiprows 0..5.
    assert plan.skiprows == [0, 1, 2, 3, 4, 5]
    assert plan.header == 0          # post-skip normalized
    assert plan.nrows == 3

    df = load_dataframe(path, plan)
    assert list(df.columns) == ["제품명", "단가", "재고", "비고"]
    assert len(df) == 3
    assert list(df["제품명"]) == ["키보드", "마우스", "모니터"]
    assert int(df["재고"].sum()) == 52


def test_block1_read_plan_pandas_round_trip(fixture_path) -> None:
    """T1's plan loads exactly 3 rows x 3 columns (guard 5 counterpart)."""

    path = fixture_path("multi_table_stacked")
    sheet = inspect(path).sheets[0]
    b1 = sheet.blocks[0]
    plan = b1.read_plan
    assert plan is not None
    assert plan.skiprows == []
    assert plan.header == 0
    assert plan.nrows == 3
    assert plan.usecols == "A:C"

    df = load_dataframe(path, plan)
    assert list(df.columns) == ["부서", "인원", "예산"]
    assert len(df) == 3
    assert int(df["예산"].sum()) == 9700


# ---------------------------------------------------------------------------
# Mirror rule — flat fields == blocks[0]; mirror plan == sheet.read_plan
# ---------------------------------------------------------------------------


def test_flat_fields_mirror_top_block(fixture_path) -> None:
    """The flat fields mirror the TOP-most table block (spec §10 intent).

    v1 picked the best-*scoring* block (here the wider/lower table); Phase 10b
    decides by position: blocks[0] is the row-1 table and the flat fields
    follow it.
    """

    sheet = inspect(fixture_path("multi_table_stacked")).sheets[0]
    assert len(sheet.blocks) == 2
    top = sheet.blocks[0]
    assert (top.band_start_row, top.band_end_row) == (1, 4)
    assert sheet.header_row == top.header_row == 1
    assert sheet.header_confidence == top.header_confidence
    assert sheet.header_provenance == top.header_provenance
    assert sheet.data_start_row == top.data_start_row == 2
    assert sheet.data_end_row == top.data_end_row == 4
    assert sheet.data_left_col == top.data_left_col
    assert sheet.data_right_col == top.data_right_col
    assert sheet.skip_rows == top.skip_rows == []
    assert sheet.columns == top.columns
    assert sheet.read_plan == top.read_plan


def test_mirror_block_plan_equals_sheet_plan_across_corpus(
    openable_fixture: Path,
) -> None:
    """blocks[0].read_plan == sheet.read_plan for every sheet with blocks.

    For single-band sheets the two plans are computed *independently* (the
    flat one through the v1 path, the block one through the per-block
    synthesizer), so this is a real compatibility guard, not an identity
    tautology.
    """

    profile = inspect(openable_fixture)
    for sheet in profile.sheets:
        if not sheet.blocks:
            continue
        top = sheet.blocks[0]
        assert top.read_plan is not None
        assert top.read_plan == sheet.read_plan
        # Flat mirror holds field-by-field too.
        assert sheet.header_row == top.header_row
        assert sheet.data_start_row == top.data_start_row
        assert sheet.data_end_row == top.data_end_row
        assert sheet.skip_rows == top.skip_rows
        assert sheet.columns == top.columns


def test_single_band_sheet_has_single_mirror_block(fixture_path) -> None:
    """offset_plus_subtotals: one band -> one mirror block, v1 fields intact."""

    sheet = inspect(fixture_path("offset_plus_subtotals")).sheets[0]
    assert len(sheet.blocks) == 1
    (block,) = sheet.blocks
    assert block.block_index == 0
    assert (block.band_start_row, block.band_end_row) == (1, 13)
    # v1 goldens unchanged (FIXTURES coordinates).
    assert sheet.header_row == block.header_row == 4
    assert (block.data_start_row, block.data_end_row) == (5, 11)
    assert block.skip_rows == [8, 12, 13]
    assert block.read_plan == sheet.read_plan
    assert block.read_plan is not sheet.read_plan  # independently computed


def test_headerless_fallback_sheet_has_no_blocks(fixture_path) -> None:
    """no_header (needs-manual): no table block -> v1 fallback path intact."""

    sheet = inspect(fixture_path("no_header")).sheets[0]
    assert sheet.blocks == []
    assert sheet.needs_manual_header is True
    assert sheet.read_plan is not None  # v1 fallback plan survives


def test_headerless_override_sheet_has_no_blocks(fixture_path) -> None:
    """An explicit headerless declaration keeps the v1 flat path (guard 4)."""

    opts = InspectionOptions(
        sheet_overrides={"Sheet1": SheetOverride(header_row=None)}
    )
    path = fixture_path("no_header")
    sheet = inspect(path, opts).sheets[0]
    assert sheet.blocks == []
    assert sheet.read_plan is not None and sheet.read_plan.header is None

    wr = extract(path, options=opts)
    (table,) = wr.tables
    assert list(table.dataframe.columns) == [
        f"col_{i}" for i in range(len(table.dataframe.columns))
    ]


# ---------------------------------------------------------------------------
# Guard 4 — absolute-coordinate overrides target only the containing block
# ---------------------------------------------------------------------------


def test_header_override_applies_only_to_containing_block(fixture_path) -> None:
    """header_row=8 (inside band 2) forces block 2 only; block 1 stays heuristic."""

    opts = InspectionOptions(
        sheet_overrides={"Sheet1": SheetOverride(header_row=8)}
    )
    path = fixture_path("multi_table_stacked")
    sheet = inspect(path, opts).sheets[0]
    assert len(sheet.blocks) == 2
    b1, b2 = sheet.blocks

    # Block 1 (band 1-4) does not contain row 8 -> untouched heuristic header.
    assert b1.header_row == 1
    assert b1.header_provenance == "heuristic"
    # Block 2 (band 7-10) contains row 8 -> manual override applied.
    assert b2.header_row == 8
    assert b2.header_provenance == "manual"
    assert b2.header_confidence == 1.0
    assert (b2.data_start_row, b2.data_end_row) == (9, 10)
    assert b2.read_plan is not None and b2.read_plan.nrows == 2

    wr = extract(path, options=opts)
    t1, t2 = wr.tables
    assert list(t1.dataframe.columns) == ["부서", "인원", "예산"]
    assert len(t1.dataframe) == 3
    # Row 8 became T2's header; only rows 9-10 remain as data.
    assert len(t2.dataframe) == 2
    assert "키보드" in [str(c) for c in t2.dataframe.columns]
    assert list(t2.dataframe.iloc[:, 0]) == ["마우스", "모니터"]


def test_skip_rows_add_applies_only_to_containing_block(fixture_path) -> None:
    """skip_rows_add=[9] (absolute 1-based) folds into block 2 only."""

    opts = InspectionOptions(
        sheet_overrides={"Sheet1": SheetOverride(skip_rows_add=[9])}
    )
    path = fixture_path("multi_table_stacked")
    sheet = inspect(path, opts).sheets[0]
    b1, b2 = sheet.blocks
    assert b1.skip_rows == []          # row 9 is outside band 1-4
    assert b2.skip_rows == [9]
    assert b2.read_plan is not None
    assert 8 in b2.read_plan.skiprows  # 1-based 9 -> 0-based 8 [D1]
    # nrows is the whole inclusive span; interior skips are not subtracted.
    assert b2.read_plan.nrows == 3

    wr = extract(path, options=opts)
    t1, t2 = wr.tables
    assert len(t1.dataframe) == 3      # block 1 untouched
    assert list(t2.dataframe["제품명"]) == ["키보드", "모니터"]  # 마우스 skipped


# ---------------------------------------------------------------------------
# Guard 7 — non-table bands enter no block, with a visible warning
# ---------------------------------------------------------------------------


def _write_table_plus_numeric_footnote(path: Path) -> None:
    """Table rows 1-4, blank rows 5-6, then a numeric-only footnote row 7."""

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["item", "qty", "price"])
    ws.append(["A-1", 10, 1.5])
    ws.append(["A-2", 20, 2.5])
    ws.append(["A-3", 30, 3.5])
    # Rows 5-6 blank (BLANK_RUN separator); row 7 numeric footnote (no header
    # signal: zero string cells -> §7.1 score 0.0 < threshold).
    ws.cell(row=7, column=1, value=99999)
    wb.save(path)


def test_non_table_band_is_skipped_with_warning(tmp_path) -> None:
    """A footnote band yields no block — a warning, never silent loss."""

    path = tmp_path / "footnote_band.xlsx"
    _write_table_plus_numeric_footnote(path)

    profile = inspect(path)
    sheet = profile.sheets[0]
    assert len(sheet.blocks) == 1          # only the real table
    assert sheet.blocks[0].header_row == 1
    assert profile.open_errors             # the skip is visible (wording unpinned)

    wr = extract(path)
    assert [t.table_id for t in wr.tables] == ["Sheet1!T1"]
    assert len(wr.tables[0].dataframe) == 3
    assert wr.warnings


# ---------------------------------------------------------------------------
# Guard 6 — deterministic warnings + JSON determinism
# ---------------------------------------------------------------------------


def test_multi_table_warnings_are_deterministic(fixture_path) -> None:
    p = fixture_path("multi_table_stacked")
    first = extract(p)
    second = extract(p)
    assert first.warnings == second.warnings
    assert first.to_json() == second.to_json()


def test_stacked_uneven_json_is_deterministic(fixture_path) -> None:
    p = fixture_path("stacked_uneven_width")
    assert extract(p).to_json() == extract(p).to_json()


# ---------------------------------------------------------------------------
# Existing goldens unchanged (Task 10.2 Step 6 spot-checks)
# ---------------------------------------------------------------------------


def test_blank_run_terminates_flat_golden_unchanged(fixture_path) -> None:
    """The v1 table block of blank_run_terminates is byte-for-byte intact.

    The noise rows 9-10 form a second band that the heuristics judge
    table-like (a leading label row over a data row), so it now surfaces as a
    characterized T2 instead of being silently absorbed — but T1 and the flat
    sheet fields keep their exact v1 golden values.
    """

    path = fixture_path("blank_run_terminates")
    sheet = inspect(path).sheets[0]
    assert sheet.header_row == 1
    assert (sheet.data_start_row, sheet.data_end_row) == (2, 5)
    assert sheet.skip_rows == []
    assert sheet.read_plan is not None and sheet.read_plan.nrows == 4

    wr = extract(path)
    t1 = wr.tables[0]
    assert t1.table_id == "Sheet1!T1"
    assert list(t1.dataframe.columns) == ["name", "qty", "price"]
    assert len(t1.dataframe) == 4
    flat = t1.dataframe.astype(str).to_numpy().ravel().tolist()
    assert "Z-9" not in flat


def test_offset_plus_subtotals_extract_golden_unchanged(fixture_path) -> None:
    """Keystone single-table golden: still exactly one table, 6 rows, sum 590."""

    wr = extract(fixture_path("offset_plus_subtotals"))
    (table,) = wr.tables
    assert table.table_id == "Sheet1!T1"
    d = table.to_dict()
    assert d["row_count"] == 6
    assert sum(r["amount"] for r in d["records"]) == 590


# ---------------------------------------------------------------------------
# Workflow-A review fixes — title/footnote bands must never become tables
# (HIGH: not-a-table judgment + band clamp; MEDIUM: threshold edge, stray
# override; LOW: warning wording, 1-column advisory). Fixture coordinates are
# documented in tests/fixtures/generate.py (FIXTURES, 1-based [D1]).
# ---------------------------------------------------------------------------


def test_title_between_tables_extracts_exactly_two_tables(fixture_path) -> None:
    """HIGH golden: the 1-row title band yields no garbage T1 frame.

    Before the fix the title band (score exactly 0.500, no data) became a
    (12, 1) frame with the title as its column name, the real tables demoted
    to T2/T3, and both tables' column A silently duplicated into T1.
    """

    wr = extract(fixture_path("title_between_tables"))
    assert [t.table_id for t in wr.tables] == ["Sheet1!T1", "Sheet1!T2"]
    t1, t2 = wr.tables
    assert t1.header_row == 4 and t2.header_row == 10  # 1-based [D1]
    assert list(t1.dataframe.columns) == ["부서", "인원", "예산"]
    assert len(t1.dataframe) == 3
    assert list(t2.dataframe.columns) == ["품목", "수량", "금액"]
    assert len(t2.dataframe) == 3
    assert int(t2.dataframe["금액"].sum()) == 1350000


def test_title_between_tables_no_duplication_or_nan_leakage(
    fixture_path,
) -> None:
    """Every lower-table value lands in exactly one TableResult; NaN rows: 0."""

    wr = extract(fixture_path("title_between_tables"))
    frames = [t.dataframe for t in wr.tables]
    # The lower table's values appear in exactly one extracted frame (no
    # silent aggregation duplication across tables).
    for needle in ("프린터", "스캐너", "복합기", "영업"):
        occurrences = sum(
            df.astype(str).to_numpy().ravel().tolist().count(needle)
            for df in frames
        )
        assert occurrences == 1, needle
    # Zero all-NaN (blank separator) row leakage into any frame.
    for df in frames:
        assert not df.isna().all(axis=1).any()
    # The title/footnote text never becomes a column name or a value.
    for df in frames:
        cells = df.astype(str).to_numpy().ravel().tolist()
        labels = [str(c) for c in df.columns]
        for junk in ("부서별 집계", "단위는 천원"):
            assert all(junk not in c for c in labels)
            assert all(junk not in c for c in cells)


def test_title_between_tables_mirror_restores_v1_header(fixture_path) -> None:
    """The flat mirror (blocks[0]) is the row-4 table — the v1 result.

    v1's whole-sheet analysis located header_row=4; the regression mirrored
    the garbage title block (header_row=1) instead.
    """

    sheet = inspect(fixture_path("title_between_tables")).sheets[0]
    assert [
        (b.band_start_row, b.band_end_row) for b in sheet.blocks
    ] == [(4, 7), (10, 13)]
    assert sheet.header_row == 4
    assert (sheet.data_start_row, sheet.data_end_row) == (5, 7)
    assert sheet.read_plan == sheet.blocks[0].read_plan


def test_title_and_footnote_bands_warn_not_a_table(fixture_path) -> None:
    """LOW #6: the not-a-table warning states the *measured* reason.

    Both 1-row string bands (title row 1, footnote row 16) are rejected for
    unresolved data — and no garbage empty table is created for either.
    """

    wr = extract(fixture_path("title_between_tables"))
    skipped = [w for w in wr.warnings if "judged not a table" in w]
    assert any("rows 1-1" in w for w in skipped)
    assert any("rows 16-16" in w for w in skipped)
    # The truthful wording: the reason actually held (data unresolved), and
    # no false "below threshold" claim for the at-threshold title band.
    for w in skipped:
        assert "no data rows resolved" in w
        assert "below threshold" not in w
    # No garbage empty table for either band: exactly 2 tables, all populated.
    assert len(wr.tables) == 2
    assert all(len(t.dataframe) > 0 for t in wr.tables)


def test_score_at_threshold_band_rejected_via_data_check(fixture_path) -> None:
    """MEDIUM #4: the score==threshold edge is absorbed by the data check.

    A 1-row all-string band scores exactly 0.500 == the default threshold, so
    the ``score < threshold`` guard alone can never reject it; the band still
    falls because a 1-row band resolves no data (W-A review HIGH rule).
    """

    from excel_inspector.analyzers.header_locator import _score_row

    title_band_rows = [["2026년 1분기 부서별 집계", None, None]]
    score = _score_row(0, title_band_rows, 1)  # band-local col_count == 1
    assert score == 0.5
    assert score == InspectionOptions().header_confidence_threshold
    assert not (score < InspectionOptions().header_confidence_threshold)

    sheet = inspect(fixture_path("title_between_tables")).sheets[0]
    assert all(b.band_start_row != 1 for b in sheet.blocks)  # band rejected


def test_header_override_outside_all_bands_warns_and_is_ignored(
    fixture_path,
) -> None:
    """MEDIUM #5: an override pointing at a blank separator row is not silent."""

    opts = InspectionOptions(
        sheet_overrides={"Sheet1": SheetOverride(header_row=5)}
    )
    path = fixture_path("multi_table_stacked")
    profile = inspect(path, opts)
    assert any(
        "header_row override 5 falls inside no detected table band; ignored"
        in w
        for w in profile.open_errors
    )
    # Both blocks keep their heuristic headers; extraction is unaffected.
    b1, b2 = profile.sheets[0].blocks
    assert b1.header_row == 1 and b1.header_provenance == "heuristic"
    assert b2.header_row == 7 and b2.header_provenance == "heuristic"

    wr = extract(path, options=opts)
    t1, t2 = wr.tables
    assert len(t1.dataframe) == 3 and len(t2.dataframe) == 3


def test_manual_override_block_plan_clamped_to_band(fixture_path) -> None:
    """HIGH #2 defense line: an unresolved block never reads past its band.

    header_row=4 (the last row of band 1) is a manual override, so the block
    is kept even though no data resolves below it — without the clamp its
    plan would be ``nrows=None`` (read to EOF) and swallow the blank
    separator plus all of table 2 into T1.
    """

    opts = InspectionOptions(
        sheet_overrides={"Sheet1": SheetOverride(header_row=4)}
    )
    path = fixture_path("multi_table_stacked")
    sheet = inspect(path, opts).sheets[0]
    b1, b2 = sheet.blocks
    assert b1.header_row == 4 and b1.header_provenance == "manual"
    assert (b1.data_start_row, b1.data_end_row) == (None, None)
    assert b1.read_plan is not None
    assert b1.read_plan.nrows == 0          # band_end(4) - header_row(4) [D1]
    assert any("clamped to the band end" in n for n in b1.read_plan.notes)
    assert b2.header_row == 7               # block 2 untouched

    wr = extract(path, options=opts)
    t1, t2 = wr.tables
    assert len(t1.dataframe) == 0           # nothing below the manual header
    # No band leakage: table 2's values never enter T1.
    flat1 = t1.dataframe.astype(str).to_numpy().ravel().tolist()
    assert "키보드" not in flat1
    assert len(t2.dataframe) == 3
    assert list(t2.dataframe["제품명"]) == ["키보드", "마우스", "모니터"]


def test_blank_run_terminates_noise_band_one_column_note(fixture_path) -> None:
    """LOW #7/#9: the 1-column noise band T2 carries the verify advisory.

    The '기타 메모' band (rows 9-10, FIXTURES) genuinely scores as a header
    over a data row, so it stays extracted — but its band-local 1-column span
    is flagged for human verification instead of passing as a silent table.
    """

    wr = extract(fixture_path("blank_run_terminates"))
    assert [t.table_id for t in wr.tables] == ["Sheet1!T1", "Sheet1!T2"]
    t1, t2 = wr.tables
    note = "1-column band — verify this is a real table"
    assert t2.dataframe.shape == (1, 1)
    assert list(t2.dataframe.columns) == ["기타 메모"]
    assert note in t2.notes
    assert note not in t1.notes             # the real 3-column table is clean


def test_extracted_band_warning_wording(fixture_path) -> None:
    """LOW #8: a successfully extracted band reads 'extracted', not 'suspected'."""

    wr = extract(fixture_path("multi_table_stacked"))
    assert any(
        "additional table block extracted as 'Sheet1!T2' (rows 7-10)" in w
        for w in wr.warnings
    )
    assert not any("suspected" in w for w in wr.warnings)

    # title_between_tables: extracted bands are renamed, the rejected footnote
    # band keeps its suspicion alongside the not-a-table judgment.
    wr2 = extract(fixture_path("title_between_tables"))
    assert any("extracted as 'Sheet1!T1' (rows 4-7)" in w for w in wr2.warnings)
    assert any(
        "extracted as 'Sheet1!T2' (rows 10-13)" in w for w in wr2.warnings
    )
    suspected = [w for w in wr2.warnings if "suspected" in w]
    assert suspected == [
        "sheet 'Sheet1': additional table block suspected at rows 16-16"
    ]


def test_title_between_tables_json_is_deterministic(fixture_path) -> None:
    p = fixture_path("title_between_tables")
    assert extract(p).to_json() == extract(p).to_json()
